"""Uploaded test case sheets: saved safely, and turned into CSV Claude can read."""

# openpyxl ships no type information.
# pyright: reportMissingTypeStubs=false, reportUnknownMemberType=false, reportUnknownVariableType=false
# pyright: reportUnknownArgumentType=false, reportOptionalMemberAccess=false

import csv
import io
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from nkqa import imports, planner
from nkqa import workspace as workspace_mod
from nkqa.workspace import Workspace


@pytest.fixture
def ws(tmp_path: Path) -> Workspace:
	return workspace_mod.create(tmp_path)


def workbook() -> bytes:
	book = Workbook()
	cases = book.active
	cases.title = 'Regression Cases'
	cases.append(['ID', 'Module', 'Title', 'Steps', 'Expected', 'Due', 'Estimate'])
	cases.append(['TC-1', 'Projects', 'Create a project', '1. Open\n2. Save', 'Saved', datetime(2026, 10, 1), '=2*3'])
	cases.append(['TC-2', None, 'Rename a project', 'Rename it', 'Renamed', None, 4.0])
	cases.append([None] * 7)  # a formatted but empty row
	cases.append(['TC-3', 'Clients', 'Find a client', 'Search', 'Found', None, None])
	cases.merge_cells('B2:B3')  # "Projects" spans TC-1 and TC-2
	book.create_sheet('Notes')  # nothing in it: no CSV for it
	smoke = book.create_sheet('Smoke')
	smoke.append(['ID', 'Title'])
	smoke.append(['S-1', 'Sign in'])
	out = io.BytesIO()
	book.save(out)
	return out.getvalue()


def rows(path: Path) -> list[list[str]]:
	return list(csv.reader(path.open(encoding='utf-8')))


def test_each_sheet_with_cases_becomes_a_csv(ws: Workspace) -> None:
	path = imports.save_upload(ws, 'Regression Suite.xlsx', workbook())
	assert path == ws.imports_dir / 'regression-suite.xlsx'
	sheets = imports.to_csv(path)
	assert [(s.name, s.rows) for s in sheets] == [('Regression Cases', 3), ('Smoke', 1)]

	got = rows(sheets[0].csv)
	assert got[0] == ['ID', 'Module', 'Title', 'Steps', 'Expected', 'Due', 'Estimate']
	assert [r[1] for r in got[1:]] == ['Projects', 'Projects', 'Clients']  # the merged cell, filled down
	assert got[1][3] == '1. Open\n2. Save'  # a multi-line cell stays one cell
	assert got[1][5] == '2026-10-01' and got[2][6] == '4'  # a date, and 4.0 read as 4
	assert len(got) == 4  # the empty row is gone

	note = imports.attachment_note(ws, path, sheets)
	assert note.startswith('[Attached: imports/regression-suite.xlsx — sheet "Regression Cases": ')
	assert 'imports/regression-suite--regression-cases.csv, 3 rows; columns: ID, Module, Title' in note
	assert 'sheet "Smoke"' in note


def test_a_csv_is_already_readable(ws: Workspace) -> None:
	path = imports.save_upload(ws, 'cases.csv', b'\xef\xbb\xbfID,Title\nC1,Login\n,\nC2,Logout\n')
	[sheet] = imports.to_csv(path)
	assert sheet.csv == path and sheet.rows == 2 and sheet.columns == ['ID', 'Title']  # BOM and blank row ignored
	[tsv] = imports.to_csv(imports.save_upload(ws, 'cases.tsv', b'ID\tTitle\nC1\tLogin\n'))
	assert tsv.columns == ['ID', 'Title'] and tsv.rows == 1


def test_uploads_are_named_by_us_and_never_overwrite(ws: Workspace) -> None:
	first = imports.save_upload(ws, '../../etc/Cases.xlsx', b'x')
	second = imports.save_upload(ws, 'cases.xlsx', b'y')
	assert first == ws.imports_dir / 'cases.xlsx' and second == ws.imports_dir / 'cases-2.xlsx'
	assert first.read_bytes() == b'x'


@pytest.mark.parametrize(
	('name', 'data', 'why'),
	[
		('cases.xls', b'x', 'save it as .xlsx'),
		('cases.pdf', b'x', 'Attach an .xlsx'),
		('cases.csv', b'', 'empty'),
		('cases.csv', b'x' * (imports.MAX_BYTES + 1), 'over'),
	],
)
def test_what_cannot_be_imported_says_why(ws: Workspace, name: str, data: bytes, why: str) -> None:
	with pytest.raises(ValueError, match=why):
		imports.save_upload(ws, name, data)


def test_an_imported_module_keeps_its_sub_folders() -> None:
	assert planner.folder('Projects/Creation Flow') == 'projects/creation-flow'
	assert planner.folder('Checkout Flow!') == 'checkout-flow'
	assert planner.folder('') == 'unnamed-test'
