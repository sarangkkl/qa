"""Test cases people already have: an uploaded sheet, made readable for Claude.

Claude reads CSV but not Excel, so an .xlsx becomes one CSV per sheet. Merged cells are filled
down: test case sheets merge the Module cell across all of its rows, and a row that does not say
its module cannot be filed in the right folder.
"""

# openpyxl ships no type information.
# pyright: reportMissingTypeStubs=false, reportUnknownMemberType=false, reportUnknownVariableType=false
# pyright: reportUnknownArgumentType=false

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from nkqa.workspace import Workspace, slugify

KINDS = ('.xlsx', '.csv', '.tsv')
MAX_BYTES = 20 * 1024 * 1024


@dataclass
class Sheet:
	name: str
	csv: Path
	rows: int  # test case rows, header excluded
	columns: list[str]


def save_upload(ws: Workspace, filename: str, data: bytes) -> Path:
	"""Into imports/, under a name of our making: never a path the caller chose, never an overwrite."""
	suffix = Path(filename).suffix.lower()
	if suffix == '.xls':
		raise ValueError('Old .xls files are not supported: save it as .xlsx or CSV and attach that.')
	if suffix not in KINDS:
		raise ValueError(f'Attach an .xlsx, .csv or .tsv file, not "{suffix or filename}".')
	if not data:
		raise ValueError('The file is empty.')
	if len(data) > MAX_BYTES:
		raise ValueError(f'The file is over {MAX_BYTES // 1024 // 1024} MB.')
	stem = slugify(Path(filename).stem)
	ws.imports_dir.mkdir(parents=True, exist_ok=True)
	path, n = ws.imports_dir / f'{stem}{suffix}', 2
	while path.exists():
		path, n = ws.imports_dir / f'{stem}-{n}{suffix}', n + 1
	path.write_bytes(data)
	return path


def cell(value: Any) -> str:
	if value is None:
		return ''
	if isinstance(value, datetime):
		return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=' ', timespec='minutes')
	if isinstance(value, date):
		return value.isoformat()
	if isinstance(value, float) and value.is_integer():
		return str(int(value))  # Excel stores every number as a float: test case 12, not 12.0
	return str(value).strip()


def tidy(rows: list[list[str]]) -> list[list[str]]:
	"""Drop rows and columns with nothing in them - sheets carry a lot of formatted emptiness."""
	rows = [r for r in rows if any(r)]
	if not rows:
		return []
	width = max(len(r) for r in rows)
	rows = [r + [''] * (width - len(r)) for r in rows]
	keep = [i for i in range(width) if any(r[i] for r in rows)]
	return [[r[i] for i in keep] for r in rows]


def to_csv(path: Path) -> list[Sheet]:
	"""One readable CSV per sheet that holds test cases. A CSV or TSV is already readable."""
	if path.suffix.lower() in ('.csv', '.tsv'):
		text = path.read_text(encoding='utf-8-sig', errors='replace')
		rows = tidy(list(csv.reader(io.StringIO(text), delimiter='\t' if path.suffix.lower() == '.tsv' else ',')))
		return [Sheet(path.stem, path, max(len(rows) - 1, 0), rows[0] if rows else [])]

	book = load_workbook(path, data_only=True)  # data_only: a formula cell gives its value
	sheets: list[Sheet] = []
	for page in book.worksheets:
		for merged in list(page.merged_cells.ranges):
			value = page.cell(merged.min_row, merged.min_col).value
			page.unmerge_cells(str(merged))
			# Down, not across: a Module merged over its rows applies to each row; a header merged
			# over two columns is still one header.
			for row in range(merged.min_row, merged.max_row + 1):
				page.cell(row, merged.min_col).value = value  # pyright: ignore[reportAttributeAccessIssue]  (unmerged above)
		rows = tidy([[cell(v) for v in row] for row in page.iter_rows(values_only=True)])
		if len(rows) < 2:
			continue  # empty, or a header with no test cases under it
		out = path.with_name(f'{path.stem}--{slugify(page.title)}.csv')
		with out.open('w', newline='', encoding='utf-8') as f:
			csv.writer(f).writerows(rows)
		sheets.append(Sheet(page.title, out, len(rows) - 1, rows[0]))
	return sheets


def attachment_note(ws: Workspace, path: Path, sheets: list[Sheet]) -> str:
	"""The line in front of the chat message: where to read, how much, and the columns to map."""

	def rel(p: Path) -> str:
		return p.relative_to(ws.root).as_posix()

	if not sheets:
		return f'[Attached: {rel(path)} — no sheet with test cases in it]'
	parts = [
		f'sheet "{s.name}": {rel(s.csv)}, {s.rows} rows; columns: {", ".join(c for c in s.columns if c)[:300]}'
		for s in sheets
	]
	return f'[Attached: {rel(path)} — {" | ".join(parts)}]'
